"""Parser del archivo .xlsx de ventas que envía Claro (corte diario).

Formato: tres hojas — DDI (líneas activadas), CARGAS (ventas cargadas) y
PORTABILIDAD (Pospago con portación efectiva) — cada una con un título en la
fila 1, una fila vacía y los encabezados en la fila 3. La fila de encabezados se
reconoce por sus columnas clave (en cualquier orden), así que tolera filas extra
arriba y cambios de orden. Desde octubre 2026 CARGAS viene sin SDS_NUMBER: cada
fila se toma como una venta cargada (`sin_sds`, con id sintético `CARGA-n`,
repetida según CANTIDAD o CARGA) y el corte sale de FECHA_DATO o FECHA_ACTUALIZACION.
Desde el 19/10/2026 las hojas se llaman ACTIVACIONES (o «ACTIVACONES»), PORTAS y
CARGAS, con los encabezados en la fila 1 y algunas columnas renombradas
(PLAN_DESCRIPCION_ORIG, TOTAL_ABONO, Venta, CARGA): las hojas se reconocen por
alias de nombre o, si no, por su contenido, y las columnas renombradas por alias.

Devuelve filas normalizadas (dicts con claves en minúsculas) con solo las
columnas que usa el análisis. Corre en subproceso aislado (ver jobs.py).
"""
from __future__ import annotations

from datetime import date, datetime
from typing import Any, Iterable

from openpyxl import load_workbook

# Hoja -> (nombre, columnas que identifican la fila de encabezados (todas deben estar), obligatoria).
# Antes se buscaba por la primera columna; el formato de octubre 2026 cambió el orden de CARGAS, así que se reconoce
# por columnas clave en cualquier posición.
HOJAS = {
    "ddi": ("DDI", ("FECHA_ACTIVACION", "TIPO_PRODUCTO"), True),
    "cargas": ("CARGAS", ("SDS_ESTADO", "TIPO_PRODUCTO"), True),
    "portabilidad": ("PORTABILIDAD", ("FECHA_ACTIVACION", "TIPO_PRODUCTO"), False),
}
# Cómo puede venir nombrada cada hoja (normalizado: mayúsculas, sin tildes ni espacios); después, por prefijo.
ALIAS_HOJAS = {
    "ddi": ("DDI", "ACTIVACIONES", "ACTIVACONES", "ACTIVADAS", "NETAS", "LINEAS ACTIVADAS"),
    "cargas": ("CARGAS", "VENTAS CARGADAS", "CARGA"),
    "portabilidad": ("PORTABILIDAD", "PORTAS", "PORTACIONES", "PORTA", "PORTADAS"),
}
PREFIJOS_HOJAS = {"ddi": ("DDI", "ACTIVA", "NETA"), "cargas": ("CARGA",), "portabilidad": ("PORTA",)}
# Columnas que cambiaron de nombre: la que usa el análisis → cómo puede venir en el archivo (la primera que esté).
ALIAS_COLUMNAS = {
    "PLAN_DESCRIPCION": ("PLAN_DESCRIPCION", "PLAN_DESCRIPCION_ORIG"),
    "TOTAL_PLAN": ("TOTAL_PLAN", "TOTAL_ABONO"),
    "CANTIDAD": ("CANTIDAD", "CARGA"),
    "VENTA": ("VENTA",),   # «Venta»: el encabezado se compara en mayúsculas
}
# Columnas de fecha: se normalizan a ISO (vienen como datetime, 'YYYY-MM-DD' o 'DD/MM/YYYY' según el formato).
FECHAS = {"fecha_activacion", "portacion_fecha", "sds_fecha_alta_venta", "sds_fecha_venta", "fecha_venta", "fecha_dato",
          "fecha_actualizacion"}

COLUMNAS = {
    "ddi": [
        "FECHA_ACTIVACION", "PERIODO_ACTIVACION", "TIPO_PRODUCTO", "PLAN_ID_ORIG", "PLAN_DESCRIPCION",
        "CAMPANIA", "PROMO_ID", "ACTIVACION_PORTACION", "PORTACION", "PORTACION_TIPO", "PORTACION_FECHA",
        "ORIGEN_PORTACION", "REGION", "DEPARTAMENTO", "CIUDAD", "CLIENTE_SEGMENTO", "SUBCANAL",
        "POS_ID", "POS_NOMBRE", "CONSUMO_DATOS", "SDS_NUMBER", "SDS_FORMULARIO", "LINEA_ORIG",
        "LINEA_ESTADO_CIERRE", "LINEA_RAZON_CIERRE", "VENTA", "TOTAL_PLAN", "TOTAL_DESCUENTO", "TOTAL_NETO",
    ],
    "cargas": [
        "PERIODO_CARGA_VENTA", "SDS_NUMBER", "SDS_FECHA_ALTA_VENTA", "SDS_FECHA_VENTA", "TIEMPO_ATENC_MIN",
        "SDS_ESTADO", "SDS_CANC_ADM", "SDS_FORMULARIO", "LINEA_ORIG", "TIPO_PRODUCTO", "PLAN_ID_ORIG",
        "PLAN_DESCRIPCION_ORIG", "CAMPANIA", "CAMPANIA_DESCRIPCION", "FECHA_ACTIVACION", "PERIODO_VENTA",
        "ORIGEN_PORTACION", "TIPO_PORT", "RIESGO_ORI", "DEPARTAMENTO_FACT", "CIUDAD_FACT",
        "VENDEDOR_LEGAJO", "VENDEDOR_NOMBRE", "VENDEDOR_APELLIDO", "POS_ID", "POS_NOMBRE", "SUBCANAL",
        "COMENTARIO", "FECHA_DATO",
        # Formato de octubre 2026 (sin SDS_NUMBER ni vendedor que cargó): lo que sí trae y sirve.
        "TIPO_NEGOCIO_ORIG", "FECHA_VENTA", "REGION", "FECHA_ACTUALIZACION", "CANTIDAD", "SUPERVISOR_NOMBRE",
        "CLIENTE_NOMBRE", "CLIENTE_APELLIDO",
    ],
    "portabilidad": [
        "FECHA_ACTIVACION", "TIPO_PRODUCTO", "PLAN_DESCRIPCION", "PORTACION_TIPO", "PORTACION_FECHA",
        "ORIGEN_PORTACION", "POS_NOMBRE", "SUBCANAL", "CONSUMO_DATOS", "SDS_NUMBER", "LINEA_ORIG",
        "LINEA_CIERRE", "LINEA_ESTADO_CIERRE", "LINEA_RAZON_CIERRE", "CIUDAD",
    ],
}


class ArchivoInvalido(ValueError):
    pass


def _norm(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v if v not in ("", "N/A") else None
    return v


_MESES_EN = {m: i for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), start=1)}
_MESES_ES = {m: i for i, m in enumerate(("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"), start=1)}


def _fecha_iso(v: Any) -> Any:
    """'2026-10-08', '08/10/2026', '14 Oct 2026 09:06', datetime → 'YYYY-MM-DD'. Lo que no se entiende queda igual."""
    if v is None or isinstance(v, (date, datetime)):
        return _norm(v)
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
                "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    partes = s.replace(",", " ").split()
    if len(partes) >= 3 and partes[0].isdigit() and partes[2].isdigit():
        mes = _MESES_EN.get(partes[1][:3].lower()) or _MESES_ES.get(partes[1][:3].lower())
        if mes:
            try:
                return date(int(partes[2]), mes, int(partes[0])).isoformat()
            except ValueError:
                pass
    return s or None


def _find_header(rows: Iterable[tuple], claves: tuple[str, ...], max_scan: int = 15) -> tuple[int, list[str]] | None:
    for i, row in enumerate(rows):
        if i >= max_scan:
            break
        header = [str(c).strip().upper() if c is not None else "" for c in row]
        if all(k in header for k in claves):
            return i, header
    return None


def _parse_sheet(ws, claves: tuple[str, ...], columnas: list[str]) -> list[dict[str, Any]]:
    rows = ws.iter_rows(values_only=True)
    found = _find_header(rows, claves)
    if not found:
        raise ArchivoInvalido(
            f"La hoja '{ws.title}' no tiene la fila de encabezados esperada (con las columnas {' y '.join(claves)})."
        )
    _, header = found
    idx: dict[str, int] = {}
    for c in columnas:
        for nombre in ALIAS_COLUMNAS.get(c, (c,)):
            if nombre in header:
                idx[c] = header.index(nombre)
                break
    faltan = [c for c in columnas if c not in idx]
    es_cargas = "SDS_ESTADO" in claves
    if "TIPO_PRODUCTO" in faltan or ("SDS_NUMBER" in faltan and not es_cargas):
        raise ArchivoInvalido(f"La hoja '{ws.title}' no tiene las columnas {', '.join(faltan[:5])}.")
    sin_sds = es_cargas and "SDS_NUMBER" in faltan
    out: list[dict[str, Any]] = []
    n = 0
    for row in rows:  # continúa después del encabezado
        if row is None or all(v is None for v in row):
            continue
        rec = {c.lower(): _norm(row[i]) if i < len(row) else None for c, i in idx.items()}
        for k in FECHAS & rec.keys():
            rec[k] = _fecha_iso(rec[k])
        if es_cargas:
            if rec.get("sds_fecha_venta") is None and rec.get("fecha_venta"):
                rec["sds_fecha_venta"] = rec["fecha_venta"]
            if rec.get("fecha_dato") is None and rec.get("fecha_actualizacion"):
                rec["fecha_dato"] = rec["fecha_actualizacion"]   # el corte del archivo, en el formato nuevo
        if sin_sds:
            # Sin SDS_NUMBER cada fila es una venta cargada igual, pero no se puede cruzar con su activación.
            if rec.get("sds_estado") is None and rec.get("tipo_producto") is None:
                continue
            cantidad = rec.get("cantidad")
            if isinstance(cantidad, str) and cantidad.strip().isdigit():
                cantidad = int(cantidad)
            veces = int(cantidad) if isinstance(cantidad, (int, float)) and 1 < cantidad <= 50 else 1
            for k in range(veces):
                n += 1
                out.append({**rec, "sds_number": f"CARGA-{n}", "sin_sds": True, "cantidad": cantidad})
            continue
        if rec.get("sds_number") is None:
            continue
        rec["sds_number"] = str(rec["sds_number"]).strip()
        out.append(rec)
    return out


def _nombre_hoja(titulo: str) -> str:
    import unicodedata

    s = unicodedata.normalize("NFKD", str(titulo or ""))
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).upper().split())


def _asignar_hojas(worksheets) -> dict[str, Any]:
    """Qué hoja del archivo es cada una (ddi, cargas, portabilidad): por nombre exacto o alias, después por prefijo
    y, si todavía falta alguna, por contenido (CARGAS es la que tiene SDS_ESTADO; de las que tienen FECHA_ACTIVACION,
    la primera es activaciones y la segunda portabilidad)."""
    libres = list(worksheets)
    out: dict[str, Any] = {}
    for key, alias in ALIAS_HOJAS.items():
        ws = next((w for w in libres if _nombre_hoja(w.title) in alias), None)
        if ws is not None:
            out[key], libres = ws, [w for w in libres if w is not ws]
    for key, prefijos in PREFIJOS_HOJAS.items():
        if key in out:
            continue
        ws = next((w for w in libres if _nombre_hoja(w.title).startswith(prefijos)), None)
        if ws is not None:
            out[key], libres = ws, [w for w in libres if w is not ws]
    if len(out) < len(HOJAS) and libres:
        def tiene(ws, claves) -> bool:
            return bool(_find_header(ws.iter_rows(values_only=True), claves))
        if "cargas" not in out:
            ws = next((w for w in libres if tiene(w, HOJAS["cargas"][1])), None)
            if ws is not None:
                out["cargas"], libres = ws, [w for w in libres if w is not ws]
        con_activacion = [w for w in libres if tiene(w, HOJAS["ddi"][1]) and not tiene(w, ("SDS_ESTADO",))]
        for key in ("ddi", "portabilidad"):
            if key not in out and con_activacion:
                out[key] = con_activacion.pop(0)
    return out


def parse_ventas_netas(path: str) -> dict[str, Any]:
    """Lee el .xlsx y devuelve {'ddi': [...], 'cargas': [...], 'portabilidad': [...], 'hojas': [...]}."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # archivo corrupto o no es xlsx
        raise ArchivoInvalido("No se pudo abrir el archivo. Debe ser un .xlsx de Excel válido.") from exc
    try:
        result: dict[str, Any] = {"hojas": [ws.title for ws in wb.worksheets]}
        asignadas = _asignar_hojas(wb.worksheets)
        for key, (nombre, claves, obligatoria) in HOJAS.items():
            ws = asignadas.get(key)
            if ws is None:
                if obligatoria:
                    raise ArchivoInvalido(
                        f"Falta la hoja '{nombre}' (o ACTIVACIONES / PORTAS). El archivo debe tener las hojas de "
                        "activaciones (DDI), CARGAS y portabilidad (PORTAS)."
                    )
                result[key] = []
                continue
            result[key] = _parse_sheet(ws, claves, COLUMNAS[key])
        if not result["ddi"] and not result["cargas"]:
            raise ArchivoInvalido("El archivo no contiene filas de datos.")
        return result
    finally:
        wb.close()
