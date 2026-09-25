"""Planilla descargable (.xlsx) del informe de Ventas Netas: visión operativa."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="0F1116")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_ALERT_FILL = PatternFill("solid", fgColor="FDE2E0")
_ESTADO_LABEL = {"Vta_Finalizada": "Finalizadas", "Vta_A_Confirmar": "A confirmar", "Vta_Procesado": "Procesadas", "Vta_Rechazada": "Rechazadas"}


def _sheet(wb: Workbook, title: str, headers: list[tuple[str, str]], rows: list[dict[str, Any]],
           alerta_key: str | None = None) -> None:
    ws = wb.create_sheet(title)
    ws.append([h for h, _ in headers])
    for c in ws[1]:
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(k) for _, k in headers])
        if alerta_key and r.get(alerta_key):
            for c in ws[ws.max_row]:
                c.fill = _ALERT_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, (h, _) in enumerate(headers, start=1):
        ancho = max(len(h), *(len(str(r.get(headers[i - 1][1]) or "")) for r in rows[:300])) if rows else len(h)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, ancho + 2), 45)


def build_xlsx(report) -> bytes:
    d = report.data or {}
    k = d.get("kpis", {})
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Ventas Netas · Televentas CLARO"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Período {report.periodo} · corte al {k.get('fecha_dato') or '—'} · estado: {report.status}"])
    ws.append([])
    for label, key in [
        ("Ventas netas", "netas"), ("Pospago", "pospago"), ("GPON", "gpon"), ("IPTV", "iptv"),
        ("Portadas", "portadas"), ("Nativas", "nativas"), ("% portación", "pct_portacion"),
        ("Pospago con uso", "pospago_con_uso"), ("Pospago sin uso", "pospago_sin_uso"), ("% sin uso", "pct_sin_uso"),
        ("Líneas suspendidas", "suspendidas"), ("Portadas fuera de netas", "fuera_de_netas"),
        ("Cargas del mes", "cargas"), ("Cargas finalizadas", "cargas_finalizadas"),
        ("Pendientes", "pendientes"), ("Pendientes de portación", "pendientes_portacion"),
        ("Pendientes con más de 7 días", "pendientes_mas_de_7_dias"),
        ("Vendedores", "vendedores"), ("Vendedores en alerta", "vendedores_alerta"),
    ]:
        ws.append([label, k.get(key)])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14

    _sheet(wb, "Vendedores", [
        ("Vendedor", "vendedor"), ("Subcanal", "subcanal"), ("Pospago total", "pospago"),
        ("Pospago SIN USO", "sin_uso"), ("Pospago con uso", "con_uso"), ("% líneas en USO", "pct_uso"),
        ("Ventas GPON", "gpon"), ("Ventas IPTV", "iptv"), ("Total netas", "total"),
        ("Portadas", "portadas"), ("Suspendidas", "suspendidas"), ("Alerta", "alerta"),
    ], d.get("vendedores", []), alerta_key="alerta")

    _sheet(wb, "Detalle netas", [
        ("SDS", "sds_number"), ("Línea", "linea"), ("Fecha activación", "fecha_activacion"),
        ("Producto", "producto"), ("Plan", "plan"), ("Campaña", "campania"), ("Portación", "portacion"),
        ("Tipo portación", "tipo_port"), ("Origen", "origen_portacion"), ("Consumo", "consumo"),
        ("Estado línea", "estado_linea"), ("Razón cierre", "razon_cierre"), ("Vendedor", "vendedor"),
        ("Subcanal", "subcanal"), ("Ciudad", "ciudad"), ("Segmento", "segmento"), ("Total neto", "total_neto"),
    ], d.get("detalle_netas", []))

    _sheet(wb, "Pendientes", [
        ("SDS", "sds_number"), ("Fecha alta", "fecha_alta"), ("Días", "dias"), ("Antigüedad", "antiguedad"),
        ("Estado", "estado"), ("Canc. adm.", "cancelacion_adm"), ("Producto", "producto"), ("Plan", "plan"),
        ("Campaña", "campania"), ("Portación", "portacion"), ("Origen", "origen_portacion"), ("Riesgo", "riesgo"),
        ("Legajo", "legajo"), ("Cargado por", "cargado_por"), ("Ciudad", "ciudad"), ("Comentario", "comentario"),
    ], d.get("pendientes", {}).get("detalle", []))

    sh = d.get("sali_hablando", {})
    _sheet(wb, "Sali hablando", [
        ("Sin uso", "sin_uso"), ("Fecha portación", "fecha_portacion"), ("Activación", "fecha_activacion"),
        ("Días act. → port.", "dias_activacion_a_portacion"), ("Días desde portación", "dias_desde_portacion"),
        ("SDS", "sds_number"), ("Línea", "linea"), ("Plan", "plan"), ("Origen", "origen_portacion"), ("Consumo", "consumo"),
        ("Estado línea", "estado_linea"), ("Razón cierre", "razon_cierre"), ("Vendedor", "vendedor"), ("Subcanal", "subcanal"),
        ("Ciudad", "ciudad"), ("En DDI", "en_ddi"), ("Riesgo carga", "riesgo"),
    ], sh.get("detalle", []), alerta_key="sin_uso")

    _sheet(wb, "Fuera de netas", [
        ("SDS", "sds_number"), ("Línea", "linea"), ("Fecha activación", "fecha_activacion"), ("Plan", "plan"),
        ("Tipo portación", "tipo_port"), ("Origen", "origen_portacion"), ("Consumo", "consumo"),
        ("Estado línea", "estado_linea"), ("Razón cierre", "razon_cierre"), ("Vendedor", "vendedor"), ("Ciudad", "ciudad"),
    ], d.get("fuera_de_netas", {}).get("detalle", []))

    prod = d.get("productividad", {})
    estados = prod.get("estados", [])
    est_cols = [(_ESTADO_LABEL.get(e, e), e) for e in estados]
    _sheet(wb, "Evolutivo diario", [
        ("Día", "dia"), ("Cargas", "total"), ("Acumulado", "acumulado"), *est_cols, ("% finalización", "pct_finalizacion"),
        ("Pospago", "pospago"), ("Internet (IF)", "internet"), ("IPTV", "iptv"),
        ("Capital y Central", "capital_central"), ("Interior", "interior"),
    ], prod.get("por_dia", []))
    _sheet(wb, "Cargas por vendedor", [
        ("Vendedor", "vendedor"), ("Subcanal", "subcanal"), ("Cargas", "total"), *est_cols, ("% finalización", "pct_finalizacion"),
        ("Pospago", "pospago"), ("Internet (IF)", "internet"), ("IPTV", "iptv"),
        ("Capital y Central", "capital_central"), ("Interior", "interior"),
        ("Pospago con uso", "con_uso"), ("Pospago SIN USO", "sin_uso"), ("% sin uso", "pct_sin_uso"), ("Sin dato de uso", "sin_dato_uso"),
        ("Riesgo A", "riesgo_A"), ("Riesgo M", "riesgo_M"), ("Riesgo B", "riesgo_B"), ("Sin uso riesgo A", "sin_uso_riesgo_A"),
        ("Atribuidas por legajo", "por_legajo"),
    ], prod.get("por_vendedor", []))
    _sheet(wb, "Zonas", [
        ("Departamento", "departamento"), ("Zona", "zona"), ("Cargas", "total"), *est_cols, ("% finalización", "pct_finalizacion"),
        ("Pospago", "pospago"), ("Internet (IF)", "internet"), ("IPTV", "iptv"),
    ], prod.get("por_departamento", []))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
