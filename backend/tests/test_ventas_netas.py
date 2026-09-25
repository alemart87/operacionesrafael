"""Ventas Netas (Televentas CLARO): parser del .xlsx de Claro, análisis y flujo de publicación
(borradores por corte, una publicación por período, reemplazo con confirmación)."""
from __future__ import annotations

import asyncio
import uuid
from datetime import date, datetime

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook

from app.core.database import Base, engine, session_scope
from app.main import app
from app.operativas.televentas_claro.ventas_netas import jobs
from app.operativas.televentas_claro.ventas_netas.analyzer import analyze_ventas_netas, vendedor_de
from app.operativas.televentas_claro.ventas_netas.models import VentasNetasReport, VentasNetasUpload
from app.operativas.televentas_claro.ventas_netas.parser import ArchivoInvalido, parse_ventas_netas

BASE = "/api/v1/televentas-claro/ventas-netas"
TC = "televentas_claro"

DDI_COLS = ["FECHA_ACTIVACION", "PERIODO_ACTIVACION", "TIPO_PRODUCTO", "PLAN_DESCRIPCION", "CAMPANIA", "PORTACION",
            "PORTACION_TIPO", "ORIGEN_PORTACION", "CIUDAD", "CLIENTE_SEGMENTO", "SUBCANAL", "POS_ID", "POS_NOMBRE",
            "CONSUMO_DATOS", "SDS_NUMBER", "LINEA_ORIG", "LINEA_ESTADO_CIERRE", "LINEA_RAZON_CIERRE", "VENTA", "TOTAL_NETO"]
CAR_COLS = ["PERIODO_CARGA_VENTA", "SDS_NUMBER", "SDS_FECHA_ALTA_VENTA", "SDS_FECHA_VENTA", "SDS_ESTADO", "SDS_CANC_ADM",
            "TIPO_PRODUCTO", "PLAN_DESCRIPCION_ORIG", "CAMPANIA_DESCRIPCION", "FECHA_ACTIVACION", "TIPO_PORT",
            "ORIGEN_PORTACION", "RIESGO_ORI", "DEPARTAMENTO_FACT", "CIUDAD_FACT", "VENDEDOR_LEGAJO", "VENDEDOR_NOMBRE", "VENDEDOR_APELLIDO",
            "POS_NOMBRE", "SUBCANAL", "COMENTARIO", "FECHA_DATO"]
POR_COLS = ["FECHA_ACTIVACION", "TIPO_PRODUCTO", "PLAN_DESCRIPCION", "PORTACION_TIPO", "ORIGEN_PORTACION", "POS_NOMBRE",
            "SUBCANAL", "CONSUMO_DATOS", "SDS_NUMBER", "LINEA_ORIG", "LINEA_ESTADO_CIERRE", "LINEA_RAZON_CIERRE", "CIUDAD"]

D = datetime
CORTE = D(2026, 9, 22)


def _ddi(sds, fecha, producto, plan, port, consumo, pos, estado="A", razon="PNPINP"):
    return [fecha, fecha.strftime("%Y%m"), producto, plan, "726TLK", port, "SI-PreSusp" if port == "SI" else "NO",
            "TIGO" if port == "SI" else None, "ASUNCION", "MASIVO", "TKM", "1", pos, consumo, sds, f"L{sds}", estado, razon, 1, 0]


def _carga(sds, alta, estado, producto="Pospago", venta=None, port="SI-PreSusp", legajo="EXP1", pos=None, riesgo="M", depto="CAPITAL"):
    return ["202609", sds, alta, venta, estado, "NO", producto, "Control 15GB", "CAMPAÑA X", venta, port,
            "TIGO" if port != "NO" else None, riesgo, depto, "ASUNCION", legajo, "NILDA", "CACERES", pos, "TKM" if pos else None, None, CORTE]


def _por(sds, fecha, tipo, consumo, pos):
    return [fecha, "Pospago", "Control 15GB", tipo, "TIGO", pos, "TKM", consumo, sds, f"L{sds}", "A", "PNPINP", "ASUNCION"]


def build_xlsx(path, *, sin_portabilidad=False, ddi_extra=(), cargas_extra=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "DDI"
    ws.append(["Detalles para Cuenta de TIPO_PRODUCTO - ENTIDAD_PADRE: VOICENTER S.A."])
    ws.append([])
    ws.append(DDI_COLS)
    filas = [
        _ddi(1001, D(2026, 9, 5), "Pospago", "Control 15GB", "SI", "SI", "TKM - ANA PEREZ"),
        _ddi(1002, D(2026, 9, 6), "Pospago", "Control 30GB", "SI", "NO", "TKM - ANA PEREZ"),
        _ddi(1003, D(2026, 9, 7), "Pospago", "Control 15GB", "NO", "SI", "TKM -  JUAN  LOPEZ"),
        _ddi(1004, D(2026, 9, 8), "Pospago", "Control 50GB", "SI", "NO", "TKM - JUAN LOPEZ", estado="S", razon="PNT"),
        _ddi(1005, D(2026, 9, 9), "GPON", "Internet 400Mbps", "NO", "NO", "TKM - ANA PEREZ"),
        _ddi(1006, D(2026, 9, 9), "IPTV", "Claro TV", "NO", "NO", "ADG - LUIS SOSA"),
        _ddi(1007, D(2026, 8, 30), "Pospago", "Control 15GB", "SI", "SI", "TKM - ANA PEREZ"),  # otro mes
        *ddi_extra,
    ]
    for f in filas:
        ws.append(f)
    ws = wb.create_sheet("CARGAS")
    ws.append(["Detalles para Cuenta de SDS_NUMBER"])
    ws.append([])
    ws.append(CAR_COLS)
    for f in [
        *[_carga(s, D(2026, 9, 5), "Vta_Finalizada", venta=D(2026, 9, 5), pos="TKM - ANA PEREZ") for s in (1001, 1002, 1003, 1004)],
        _carga(1005, D(2026, 9, 9), "Vta_Finalizada", producto="GPON", venta=D(2026, 9, 9), port="NO", pos="TKM - ANA PEREZ"),
        _carga(1006, D(2026, 9, 9), "Vta_Finalizada", producto="IPTV", venta=D(2026, 9, 9), port="NO", legajo="EXP3", pos="ADG - LUIS SOSA"),
        _carga(1009, D(2026, 9, 10), "Vta_Finalizada", venta=D(2026, 9, 10)),  # finalizada sin activar
        _carga(2001, D(2026, 9, 21), "Vta_A_Confirmar"),                          # 1 día
        _carga(2002, D(2026, 9, 12), "Vta_A_Confirmar", legajo="EXP2"),           # 10 días
        _carga(2003, D(2026, 9, 1), "Vta_Procesado"),                             # 21 días
        _carga(2004, D(2026, 9, 18), "Vta_Rechazada", port="NO", depto="ALTO PARANA"),  # 4 días, nativa, Interior
        *cargas_extra,
    ]:
        ws.append(f)
    if not sin_portabilidad:
        ws = wb.create_sheet("PORTABILIDAD")
        ws.append(["Detalles"])
        ws.append([])
        ws.append(POR_COLS)
        ws.append(_por(1001, D(2026, 9, 5), "SI-PreSusp", "SI", "TKM - ANA PEREZ"))
        ws.append(_por(3001, D(2026, 9, 11), "SI-SaliHbl", "NO", "TKM - ANA PEREZ"))  # fuera de netas
    wb.save(path)
    return str(path)


@pytest.fixture(scope="module")
def xlsx(tmp_path_factory):
    return build_xlsx(tmp_path_factory.mktemp("vn") / "Ventas_Sep.xlsx")


def test_vendedor_normalizado():
    assert vendedor_de("TKM -  DOLLY  GONZALEZ", None) == ("DOLLY GONZALEZ", "TKM")
    assert vendedor_de("ADG - LILIAM TORRES", "ADG") == ("LILIAM TORRES", "ADG")
    assert vendedor_de(None, "TKM") == ("SIN VENDEDOR", "TKM")


def test_parser_y_analisis(xlsx):
    parsed = parse_ventas_netas(xlsx)
    assert {k: len(v) for k, v in parsed.items() if isinstance(v, list)} == {"hojas": 3, "ddi": 7, "cargas": 11, "portabilidad": 2}
    assert parsed["ddi"][0]["fecha_activacion"] == "2026-09-05"  # fechas normalizadas a ISO

    a = analyze_ventas_netas(parsed)
    k = a["kpis"]
    assert k["periodo"] == "2026-09" and k["fecha_dato"] == "2026-09-22"
    assert (k["netas"], k["pospago"], k["gpon"], k["iptv"]) == (6, 4, 1, 1)
    assert k["fuera_periodo"] == 1  # la activación de agosto no entra
    assert (k["portadas"], k["nativas"]) == (3, 3)
    assert (k["pospago_sin_uso"], k["pospago_con_uso"], k["pct_sin_uso"]) == (2, 2, 50.0)
    assert k["suspendidas"] == 1 and a["suspendidas"]["por_razon"] == [{"razon": "PNT", "total": 1}]
    assert k["fuera_de_netas"] == 1 and a["fuera_de_netas"]["sin_uso"] == 1
    assert k["finalizadas_sin_activar"] == 1 and a["finalizadas_sin_activar"][0]["sds_number"] == "1009"

    # Vendedor = POS_NOMBRE sin prefijo; "JUAN LOPEZ" con espacios dobles es el mismo vendedor.
    v = {x["vendedor"]: x for x in a["vendedores"]}
    assert set(v) == {"ANA PEREZ", "JUAN LOPEZ", "LUIS SOSA"}
    assert (v["ANA PEREZ"]["pospago"], v["ANA PEREZ"]["sin_uso"], v["ANA PEREZ"]["gpon"], v["ANA PEREZ"]["pct_uso"]) == (2, 1, 1, 50.0)
    assert (v["JUAN LOPEZ"]["pospago"], v["JUAN LOPEZ"]["suspendidas"]) == (2, 1)
    assert v["LUIS SOSA"]["subcanal"] == "ADG" and v["LUIS SOSA"]["iptv"] == 1
    assert a["alertas"] == []  # nadie llega a las 5 líneas mínimas

    # Pendientes: todo lo no finalizado, con antigüedad desde el corte.
    p = a["pendientes"]
    assert (p["total"], p["portacion"], p["mas_de_7_dias"]) == (4, 3, 2)
    assert p["detalle"][0]["sds_number"] == "2003" and p["detalle"][0]["dias"] == 21
    assert {f["rango"]: f["total"] for f in p["por_antiguedad"]} == {"0-2 días": 1, "3-7 días": 1, "8-15 días": 1, "más de 15 días": 1}
    assert p["por_legajo"][0] == {"legajo": "EXP1", "cargado_por": "NILDA CACERES", "total": 3, "mas_de_7_dias": 1}
    assert a["detalle_netas"][4]["consumo"] is None  # GPON no tiene consumo

    # Productividad (CARGAS): evolutivo por fecha de alta, estados, Pospago/Internet y zonas.
    pr = a["productividad"]
    pk = pr["kpis"]
    assert (pk["cargas"], pk["finalizadas"], pk["pct_finalizacion"]) == (11, 7, 63.6)
    assert (pk["pospago"], pk["internet"], pk["iptv"]) == (9, 1, 1)
    assert (pk["capital_central"], pk["interior"]) == (10, 1)
    assert [(z["zona"], z["total"]) for z in pr["por_zona"]] == [("Capital y Central", 10), ("Interior", 1)]
    assert pr["por_departamento"][1] == {**pr["por_departamento"][1], "departamento": "ALTO PARANA", "zona": "Interior", "Vta_Rechazada": 1}
    assert pk["mejor_dia"] == "2026-09-05" and pk["mejor_dia_total"] == 4
    assert [(f["dia"], f["total"], f["acumulado"]) for f in pr["por_dia"]][:2] == [("2026-09-01", 1, 1), ("2026-09-05", 4, 5)]
    assert pr["por_estado"][0] == {"estado": "Vta_Finalizada", "total": 7, "pct": 63.6}
    # Las pendientes sin POS se atribuyen al vendedor único del legajo (EXP1 -> ANA PEREZ); EXP2 no cargó con POS.
    pv = {v["vendedor"]: v for v in pr["por_vendedor"]}
    assert pv["ANA PEREZ"]["total"] == 9 and pv["ANA PEREZ"]["por_legajo"] == 4 and pv["ANA PEREZ"]["Vta_A_Confirmar"] == 1
    assert pv["CARGADO POR EXP2 NILDA CACERES"]["total"] == 1 and pk["sin_atribuir"] == 1


def test_alerta_por_vendedor(tmp_path):
    extra = [_ddi(5000 + i, D(2026, 9, 10), "Pospago", "Control 15GB", "SI", "SI" if i < 2 else "NO", "TKM - MAL USO")
             for i in range(6)]
    a = analyze_ventas_netas(parse_ventas_netas(build_xlsx(tmp_path / "a.xlsx", ddi_extra=extra)))
    assert [(x["vendedor"], x["pospago"], x["pct_uso"]) for x in a["alertas"]] == [("MAL USO", 6, 33.3)]
    assert a["kpis"]["vendedores_alerta"] == 1


def test_parser_valida_el_archivo(tmp_path):
    a = analyze_ventas_netas(parse_ventas_netas(build_xlsx(tmp_path / "sin-por.xlsx", sin_portabilidad=True)))
    assert a["kpis"]["fuera_de_netas"] == 0  # PORTABILIDAD es opcional

    wb = Workbook()
    wb.active.title = "OTRA"
    wb.save(tmp_path / "mal.xlsx")
    with pytest.raises(ArchivoInvalido, match="Falta la hoja 'DDI'"):
        parse_ventas_netas(str(tmp_path / "mal.xlsx"))
    (tmp_path / "txt.xlsx").write_bytes(b"no es excel")
    with pytest.raises(ArchivoInvalido, match="No se pudo abrir"):
        parse_ventas_netas(str(tmp_path / "txt.xlsx"))


# ============================ API: flujo de publicación ============================
async def _ensure_schema() -> None:
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
        await conn.run_sync(Base.metadata.create_all)


def setup_module(module):
    from app.main import _seed_profiles  # sin lifespan (AsyncClient) hay que sembrar los perfiles a mano

    async def _prep():
        await _ensure_schema()
        await _seed_profiles()
    asyncio.run(_prep())


async def _login(ac, email, pwd):
    r = await ac.post("/api/v1/auth/login", json={"email": email, "password": pwd})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


async def _new_user(ac, admin, role):
    email = f"{role}-{uuid.uuid4().hex[:6]}@voicenter.com.py"
    r = await ac.post("/api/v1/users", headers=admin, json={
        "email": email, "password": "Clave1234!", "full_name": f"Usuario {role}", "role": role, "operativas": [TC]})
    assert r.status_code == 201, r.text
    return await _login(ac, email, "Clave1234!")


async def _subir_y_procesar(ac, h, xlsx, monkeypatch) -> str:
    """Sube el archivo por la API y corre el runner (sin worker ni subproceso). Devuelve el id del informe."""
    with open(xlsx, "rb") as f:
        r = await ac.post(f"{BASE}/uploads", headers=h, files={"file": ("Ventas_Sep.xlsx", f, "application/octet-stream")})
    assert r.status_code == 202, r.text
    uid = r.json()["id"]
    async with session_scope() as db:
        (await db.get(VentasNetasUpload, uid)).status = "processing"
        await db.commit()

    async def inline(func, *args):
        return func(*args)
    monkeypatch.setattr(jobs, "run_isolated", inline)
    await jobs.run_ventas_netas(uid)
    async with session_scope() as db:
        up = await db.get(VentasNetasUpload, uid)
        assert up.status == "completed", up.last_error
        rep = (await db.execute(
            __import__("sqlalchemy").select(VentasNetasReport).where(VentasNetasReport.upload_id == uid))).scalar_one()
        return rep.id


@pytest.mark.asyncio
async def test_flujo_publicacion(xlsx, monkeypatch, tmp_path):
    # Perfiles sembrados: la perfil sembrada Analista tiene gestión; Supervisor solo ve informes.
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        admin = await _login(ac, "admin@voicenter.com.py", "Test1234!")
        analista = await _new_user(ac, admin, "analista")
        supervisor = await _new_user(ac, admin, "supervisor")

        # Solo gestión sube; solo .xlsx.
        assert (await ac.post(f"{BASE}/uploads", headers=supervisor, files={"file": ("v.xlsx", b"x")})).status_code == 403
        assert (await ac.post(f"{BASE}/uploads", headers=analista, files={"file": ("v.csv", b"x")})).status_code == 400
        assert (await ac.get(f"{BASE}/uploads", headers=supervisor)).status_code == 403

        r1 = await _subir_y_procesar(ac, analista, xlsx, monkeypatch)

        # Borrador: gestión lo ve, el resto no.
        lst = (await ac.get(f"{BASE}/reports", headers=analista)).json()
        assert [(x["id"], x["status"], x["periodo"], x["netas"]) for x in lst["items"]] == [(r1, "draft", "2026-09", 6)]
        assert lst["usuarios"][lst["items"][0]["generated_by"]] == "Usuario analista"
        assert (await ac.get(f"{BASE}/reports", headers=supervisor)).json()["items"] == []
        assert (await ac.get(f"{BASE}/reports/{r1}", headers=supervisor)).status_code == 404
        assert (await ac.get(f"{BASE}/reports/{r1}/export.xlsx", headers=supervisor)).status_code == 404

        # Publicar el primero: sin conflicto.
        r = await ac.post(f"{BASE}/reports/{r1}/publish", headers=analista, json={})
        assert r.status_code == 200 and r.json()["status"] == "published"
        assert [x["id"] for x in (await ac.get(f"{BASE}/reports", headers=supervisor)).json()["items"]] == [r1]
        det = (await ac.get(f"{BASE}/reports/{r1}", headers=supervisor)).json()
        assert det["data"]["kpis"]["pospago_sin_uso"] == 2
        x = await ac.get(f"{BASE}/reports/{r1}/export.xlsx", headers=supervisor)
        assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert 'ventas-netas_2026-09_corte-2026-09-22.xlsx' in x.headers["content-disposition"]

        # Un corte más nuevo del mismo mes: segundo borrador. Publicarlo exige confirmación.
        r2 = await _subir_y_procesar(ac, analista, build_xlsx(tmp_path / "corte2.xlsx"), monkeypatch)
        r = await ac.post(f"{BASE}/reports/{r2}/publish", headers=analista, json={})
        assert r.status_code == 409, r.text
        assert r.json()["detail"]["code"] == "replace_required"
        assert r.json()["detail"]["existing"]["id"] == r1
        assert r.json()["detail"]["existing"]["published_by"] == "Usuario analista"
        estados = {x["id"]: x["status"] for x in (await ac.get(f"{BASE}/reports", headers=analista)).json()["items"]}
        assert estados == {r1: "published", r2: "draft"}  # nada cambió

        r = await ac.post(f"{BASE}/reports/{r2}/publish", headers=analista, json={"confirm_replace": True})
        assert r.status_code == 200 and r.json()["status"] == "published"
        items = {x["id"]: x for x in (await ac.get(f"{BASE}/reports", headers=analista)).json()["items"]}
        assert items[r1]["status"] == "replaced" and items[r1]["replaced_by_report_id"] == r2
        assert [x["id"] for x in (await ac.get(f"{BASE}/reports", headers=supervisor)).json()["items"]] == [r2]

        # No se elimina un publicado; despublicar lo vuelve borrador y ahí sí.
        assert (await ac.delete(f"{BASE}/reports/{r2}", headers=analista)).status_code == 400
        assert (await ac.post(f"{BASE}/reports/{r2}/unpublish", headers=analista)).json()["status"] == "draft"
        assert (await ac.get(f"{BASE}/reports", headers=supervisor)).json()["items"] == []
        assert (await ac.delete(f"{BASE}/reports/{r2}", headers=supervisor)).status_code == 403
        assert (await ac.delete(f"{BASE}/reports/{r2}", headers=analista)).status_code == 200
        assert (await ac.get(f"{BASE}/reports/{r2}", headers=analista)).status_code == 404

        acciones = [a["action"] for a in (await ac.get("/api/v1/audit", headers=admin)).json()]
        for esperada in ("create_ventas_netas_upload", "publish_ventas_netas_report", "replace_ventas_netas_report",
                         "unpublish_ventas_netas_report", "delete_ventas_netas_report", "export_ventas_netas_report"):
            assert esperada in acciones, esperada
